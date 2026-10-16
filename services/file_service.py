from __future__ import annotations

import asyncio
import logging
import shutil
import zipfile
from pathlib import Path

from config import BASE_DIR, MAX_FILE_SIZE, MAX_STUDY_TEXT_CHARS, SUPPORTED_EXTENSIONS, TEMP_DIR, UPLOADS_DIR
from utils.security import sanitize_filename, validate_file

logger = logging.getLogger(__name__)


def _save_uploaded_file_sync(file_bytes: bytes, filename: str, user_id: int) -> str:
    """Save a Telegram upload safely inside the bot's uploads directory."""
    if not isinstance(file_bytes, (bytes, bytearray)):
        raise TypeError("بيانات الملف غير صالحة.")

    if len(file_bytes) > MAX_FILE_SIZE:
        raise ValueError(
            f"حجم الملف {len(file_bytes)/1024/1024:.1f}MB "
            f"ويتجاوز الحد المسموح {MAX_FILE_SIZE/1024/1024:.0f}MB."
        )

    name = sanitize_filename(filename)
    ext = Path(name).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError("نوع الملف غير مدعوم. المسموح: PDF/DOCX/PPTX/TXT.")

    directory = UPLOADS_DIR / str(user_id)
    directory.mkdir(parents=True, exist_ok=True)

    path = directory / name
    i = 1
    while path.exists():
        path = directory / f"{Path(name).stem}_{i}{Path(name).suffix}"
        i += 1

    path.write_bytes(bytes(file_bytes))
    return str(path)


async def save_uploaded_file(file_bytes: bytes, filename: str, user_id: int) -> str:
    return await asyncio.to_thread(_save_uploaded_file_sync, file_bytes, filename, user_id)


def _extract_text_sync(file_path: str) -> str:
    """Extract text from supported academic document formats."""
    ok, msg = validate_file(file_path)
    if not ok:
        raise ValueError(msg)

    p = Path(file_path)
    ext = p.suffix.lower()
    if ext in {'.docx', '.pptx'}:
        try:
            with zipfile.ZipFile(p) as archive:
                infos = archive.infolist()
                uncompressed = sum(int(i.file_size) for i in infos)
                if len(infos) > 5000 or uncompressed > 200 * 1024 * 1024:
                    raise ValueError('الملف المضغوط غير آمن أو حجمه الداخلي كبير جدًا.')
        except zipfile.BadZipFile as exc:
            raise ValueError('الملف ليس ملف DOCX/PPTX صالحًا.') from exc

    if ext == ".pdf":
        reader = None
        last_error = None

        # Prefer pypdf when installed, then fall back to PyPDF2.
        for module_name in ("pypdf", "PyPDF2"):
            try:
                if module_name == "pypdf":
                    from pypdf import PdfReader
                else:
                    from PyPDF2 import PdfReader
                reader = PdfReader(str(p))
                break
            except Exception as exc:
                last_error = exc

        if reader is None:
            raise RuntimeError(
                "تعذر فتح ملف PDF لأن مكتبة قراءة PDF غير مثبتة. "
                "ثبّت المتطلبات عبر: pip install -r requirements.txt"
                if last_error and "No module named" in str(last_error)
                else f"تعذر فتح ملف PDF: {last_error}"
            )

        try:
            if getattr(reader, "is_encrypted", False):
                try:
                    reader.decrypt("")
                except Exception:
                    raise ValueError("ملف PDF مشفّر بكلمة مرور ولا يمكن قراءته.")

            texts = []
            total_chars = 0
            for index, page in enumerate(reader.pages, start=1):
                try:
                    page_text = (page.extract_text() or "").strip()
                except Exception as exc:
                    logger.warning("PDF page extraction failed page=%s: %s", index, exc)
                    continue
                if not page_text:
                    continue
                total_chars += len(page_text)
                if total_chars > MAX_STUDY_TEXT_CHARS:
                    raise ValueError(
                        f"النص المستخرج من PDF يتجاوز حد المادة المحفوظة ({MAX_STUDY_TEXT_CHARS:,} حرف). "
                        "قسّم الكتاب إلى أجزاء أصغر حتى لا يتم قص المحتوى."
                    )
                texts.append(f"--- الصفحة {index} ---\n{page_text}")

            text = "\n\n".join(texts).strip()
            if not text:
                raise ValueError(
                    "تم فتح ملف PDF، لكن لم أجد نصًا قابلًا للاستخراج. "
                    "إذا كان الملف صورًا ممسوحة ضوئيًا، يحتاج OCR."
                )
            return text

        except ValueError:
            raise
        except Exception as exc:
            raise RuntimeError(f"تعذر استخراج نص PDF: {exc}") from exc

    if ext == ".docx":
        try:
            from docx import Document
            doc = Document(str(p))
            parts = [x.text for x in doc.paragraphs if x.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = "\n".join(parts).strip()
            if not text:
                raise ValueError("ملف DOCX لا يحتوي نصًا قابلًا للاستخراج.")
            return text
        except ValueError:
            raise
        except Exception as exc:
            raise RuntimeError(f"تعذر قراءة ملف DOCX: {exc}") from exc

    if ext == ".pptx":
        try:
            from pptx import Presentation
            prs = Presentation(str(p))
            parts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        parts.append(shape.text.strip())
            text = "\n".join(parts).strip()
            if not text:
                raise ValueError("ملف PPTX لا يحتوي نصًا قابلًا للاستخراج.")
            return text
        except ValueError:
            raise
        except Exception as exc:
            raise RuntimeError(f"تعذر قراءة ملف PPTX: {exc}") from exc

    try:
        text = p.read_text(encoding="utf-8", errors="replace").strip()
        if not text:
            raise ValueError("الملف النصي فارغ.")
        return text
    except Exception as exc:
        raise RuntimeError(f"تعذر قراءة الملف النصي: {exc}") from exc


async def cleanup_temp_files(max_age_hours: int = 24):
    import time
    now = time.time()
    for p in TEMP_DIR.glob("*"):
        try:
            if now - p.stat().st_mtime > max_age_hours * 3600:
                shutil.rmtree(p) if p.is_dir() else p.unlink()
        except Exception:
            pass


async def extract_text(file_path: str) -> str:
    """Run CPU/file parsing off the Telegram event loop."""
    return await asyncio.to_thread(_extract_text_sync, file_path)
